from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    HRFlowable,
    Table,
    TableStyle,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def generate_pdf(state: dict[str, Any], session_id: str, output_path: Path) -> Path:
    """Generate a clean, styled PDF version of the Business Requirements Document."""
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()
    
    # Custom styles
    primary_color = colors.HexColor("#6b21a8")
    dark_neutral = colors.HexColor("#1e1b4b")
    text_color = colors.HexColor("#334155")
    accent_color = colors.HexColor("#9333ea")
    bg_tint = colors.HexColor("#f3e8ff")

    title_style = ParagraphStyle(
        "BRDTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=24,
        leading=28,
        textColor=primary_color,
        alignment=0,
        spaceAfter=6,
    )

    meta_style = ParagraphStyle(
        "BRDMeta",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=14,
    )

    h1_style = ParagraphStyle(
        "BRDH1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=dark_neutral,
        spaceBefore=14,
        spaceAfter=8,
        keepWithNext=True,
    )

    body_style = ParagraphStyle(
        "BRDBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        textColor=text_color,
        spaceAfter=10,
    )

    finding_style = ParagraphStyle(
        "BRDFinding",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=4,
    )

    elements = []

    # Title
    elements.append(Paragraph(f"Orbius BRD — Session: {session_id.capitalize()}", title_style))
    version = state.get("version", 1)
    elements.append(Paragraph(f"Document Version: {version}.0 | Generated via Orbius Multi-Agent Studio", meta_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=primary_color, spaceBefore=0, spaceAfter=14))

    # Sections
    sections = state.get("sections", [])
    findings = state.get("findings", [])
    conflicts = state.get("conflicts", [])

    for section in sections:
        elements.append(Paragraph(section.get("title", "Untitled Section"), h1_style))
        
        sec_id = section.get("id")
        if sec_id == "open_conflicts" and conflicts:
            for c in conflicts:
                status_str = "Resolved" if c.get("status") == "resolved" else "Pending Resolution"
                topic = c.get("topic", "Conflict")
                disagreement = c.get("disagreement", "")
                
                conflict_text = f"<b>{status_str} — {topic}</b><br/>{disagreement}"
                elements.append(Paragraph(conflict_text, body_style))
                
                left_stmt = c.get("left", {}).get("statement", "")
                right_stmt = c.get("right", {}).get("statement", "")
                
                table_data = [
                    [Paragraph("<b>Source A</b>", body_style), Paragraph("<b>Source B</b>", body_style)],
                    [Paragraph(left_stmt, finding_style), Paragraph(right_stmt, finding_style)]
                ]
                t = Table(table_data, colWidths=[240, 240])
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,-1), bg_tint),
                    ('TEXTCOLOR', (0,0), (-1,-1), text_color),
                    ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#c084fc")),
                    ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#a855f7")),
                    ('TOPPADDING', (0,0), (-1,-1), 6),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                ]))
                elements.append(t)
                elements.append(Spacer(1, 10))
        else:
            # Check for specific section findings
            sec_finding_ids = section.get("finding_ids", [])
            sec_findings = [f for f in findings if f.get("id") in sec_finding_ids]
            
            if sec_findings:
                for f in sec_findings:
                    stmt = f.get("statement", "")
                    elements.append(Paragraph(f"• {stmt}", finding_style))
                elements.append(Spacer(1, 8))
            elif section.get("body"):
                elements.append(Paragraph(section["body"].replace("\n", "<br/>"), body_style))
                elements.append(Spacer(1, 6))

    doc.build(elements)
    return output_path


def generate_docx(state: dict[str, Any], session_id: str, output_path: Path) -> Path:
    """Generate a cleanly formatted Word (.docx) Business Requirements Document."""
    doc = Document()

    # Document Title
    p_title = doc.add_paragraph()
    run_title = p_title.add_run(f"Orbius BRD — Session: {session_id.capitalize()}")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(107, 33, 168)

    # Subtitle
    version = state.get("version", 1)
    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run(f"Document Version: {version}.0 | Generated via Orbius Multi-Agent Studio")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    sections = state.get("sections", [])
    findings = state.get("findings", [])
    conflicts = state.get("conflicts", [])

    for section in sections:
        h = doc.add_heading(level=1)
        h_run = h.add_run(section.get("title", "Untitled Section"))
        h_run.font.name = "Arial"
        h_run.font.size = Pt(16)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(30, 27, 75)

        sec_id = section.get("id")
        if sec_id == "open_conflicts" and conflicts:
            for c in conflicts:
                status_str = "Resolved" if c.get("status") == "resolved" else "Pending Resolution"
                p_c = doc.add_paragraph()
                p_c.add_run(f"{status_str} — {c.get('topic', 'Conflict')}: ").bold = True
                p_c.add_run(c.get("disagreement", ""))

                table = doc.add_table(rows=2, cols=2)
                table.style = 'Table Grid'
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = 'Source A'
                hdr_cells[1].text = 'Source B'
                hdr_cells[0].paragraphs[0].runs[0].font.bold = True
                hdr_cells[1].paragraphs[0].runs[0].font.bold = True

                row_cells = table.rows[1].cells
                row_cells[0].text = c.get("left", {}).get("statement", "")
                row_cells[1].text = c.get("right", {}).get("statement", "")

                doc.add_paragraph().paragraph_format.space_after = Pt(6)
        else:
            sec_finding_ids = section.get("finding_ids", [])
            sec_findings = [f for f in findings if f.get("id") in sec_finding_ids]

            if sec_findings:
                for f in sec_findings:
                    doc.add_paragraph(f.get("statement", ""), style='List Bullet')
            elif section.get("body"):
                doc.add_paragraph(section["body"])

    doc.save(str(output_path))
    return output_path
