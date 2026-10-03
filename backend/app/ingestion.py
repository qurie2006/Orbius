from __future__ import annotations

import csv
import io
import uuid
from pathlib import Path

from app.schemas import SourceChunk

TEXT_EXT = {".txt", ".md", ".eml", ".csv"}
PDF_EXT = {".pdf"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".svg"}
SHEET_EXT = {".xlsx", ".xls"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".webm", ".weba"}


def _chunks_from_text(source_id: str, file_name: str, text: str, modality: str) -> list[SourceChunk]:
    paragraphs = [p.strip() for p in text.replace("\r\n", "\n").split("\n\n") if p.strip()]
    if not paragraphs:
        paragraphs = [text.strip()] if text.strip() else []
    out: list[SourceChunk] = []
    for i, para in enumerate(paragraphs, start=1):
        out.append(
            SourceChunk(
                source_id=f"{source_id}#p{i}",
                file_name=file_name,
                modality=modality,  # type: ignore[arg-type]
                page_or_ts=f"chunk {i}",
                text=para[:4000],
            )
        )
    return out or [
        SourceChunk(
            source_id=source_id,
            file_name=file_name,
            modality=modality,  # type: ignore[arg-type]
            page_or_ts="full",
            text=text[:4000],
        )
    ]


DOCX_EXT = {".docx"}


def _extract_docx(path: Path) -> str:
    import zipfile
    import xml.etree.ElementTree as ET

    try:
        with zipfile.ZipFile(str(path)) as docx:
            xml_content = docx.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            # Find all paragraph elements w:p
            paragraphs = []
            for node in tree.iter():
                if node.tag.endswith("p"):
                    texts = [elem.text for elem in node.iter() if elem.tag.endswith("t") and elem.text]
                    if texts:
                        paragraphs.append("".join(texts))
            return "\n\n".join(paragraphs)
    except Exception:
        return ""


def ingest_file(path: Path, original_name: str | None = None) -> list[SourceChunk]:
    name = Path(original_name or path.name).name
    ext = Path(name).suffix.lower()
    source_id = f"{Path(name).stem}_{uuid.uuid4().hex[:6]}"

    try:
        if ext in IMAGE_EXT:
            return [
                SourceChunk(
                    source_id=source_id,
                    file_name=name,
                    modality="image",
                    page_or_ts="image",
                    text=f"[Image uploaded: {name}]",
                    bbox=[0.08, 0.12, 0.84, 0.62],
                )
            ]

        if ext in AUDIO_EXT:
            try:
                import os
                from app.config import settings
                api_key = os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")
                if not api_key:
                    return _chunks_from_text(source_id, name, f"[Audio uploaded: {name}. Transcription simulated because API key is missing.]", "audio")
                
                from google import genai
                client = genai.Client(api_key=api_key)
                audio_file = client.files.upload(file=str(path))
                response = client.models.generate_content(
                    model=settings.primary_model,
                    contents=[audio_file, "Generate a complete, highly detailed transcription of this audio file."]
                )
                text = getattr(response, "text", None) or f"[Audio uploaded: {name} (Empty transcription)]"
                return _chunks_from_text(source_id, name, text, "audio")
            except Exception as e:
                return _chunks_from_text(source_id, name, f"[Audio uploaded: {name}. Error during transcription: {str(e)}]", "audio")

        if ext in PDF_EXT:
            try:
                from pypdf import PdfReader

                reader = PdfReader(str(path))
                pages = []
                for i, page in enumerate(reader.pages, start=1):
                    pages.append((i, page.extract_text() or ""))
                chunks: list[SourceChunk] = []
                for i, text in pages:
                    chunks.extend(
                        _chunks_from_text(f"{source_id}_pg{i}", name, text or f"[PDF page {i}]", "pdf")
                    )
                if chunks:
                    return chunks
            except Exception:
                pass
            return _chunks_from_text(source_id, name, path.read_text(encoding="utf-8", errors="ignore"), "pdf")

        if ext in DOCX_EXT:
            docx_text = _extract_docx(path)
            if docx_text:
                return _chunks_from_text(source_id, name, docx_text, "text")

        if ext in SHEET_EXT:
            try:
                from openpyxl import load_workbook

                wb = load_workbook(path, data_only=True)
                lines = []
                for sheet in wb.worksheets:
                    lines.append(f"Sheet: {sheet.title}")
                    for row in sheet.iter_rows(max_row=50, values_only=True):
                        vals = [str(c) for c in row if c is not None]
                        if vals:
                            lines.append(" | ".join(vals))
                return _chunks_from_text(source_id, name, "\n".join(lines), "spreadsheet")
            except Exception:
                return _chunks_from_text(source_id, name, "[spreadsheet]", "spreadsheet")

        if ext == ".csv":
            text = path.read_text(encoding="utf-8", errors="ignore")
            return _chunks_from_text(source_id, name, text, "spreadsheet")

        text = path.read_text(encoding="utf-8", errors="ignore")
        return _chunks_from_text(source_id, name, text, "text")
    except Exception as e:
        return _chunks_from_text(source_id, name, f"[{name} - Error reading file content: {str(e)}]", "text")


def ingest_bytes(filename: str, data: bytes, dest: Path) -> list[SourceChunk]:
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(data)
    return ingest_file(dest, filename)
